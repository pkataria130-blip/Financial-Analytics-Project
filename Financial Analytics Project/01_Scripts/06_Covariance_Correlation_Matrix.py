#!/usr/bin/env python3
r"""
===============================================================================
Financial Analytics using Python
Script 06 : Covariance and Correlation Matrix
Author    : Vipul Mehta
===============================================================================

LEARNING OBJECTIVES
-------------------
Convert Adjusted Close prices into daily SIMPLE returns, then compute the
covariance matrix and the correlation matrix, and read them for diversification.

    Covariance  : how two assets move together (scale-dependent).
    Correlation : covariance standardized to lie in [-1, +1].
        +1 perfect positive | 0 no linear relation | -1 perfect negative.

Low / negative correlation between holdings is what makes diversification reduce
portfolio risk -- this matrix is the direct input to Script 08 (portfolio risk).

INPUT  : Yahoo_Daily_Prices_2016_2026.csv
OUTPUT : Covariance_Correlation_Matrix.xlsx
===============================================================================
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# =============================================================================
# SETTINGS
# =============================================================================
FOLDER = r"D:\Trimester 4\Finance\Financial Analytics\Financial Analytics Project\02_Output"

INPUT_FILE  = os.path.join(FOLDER, "Yahoo_Daily_Prices_2016_2026.csv")
OUTPUT_FILE = os.path.join(FOLDER, "Covariance_Correlation_Matrix.xlsx")

# =============================================================================
# STEP 1 : READ DATA
# =============================================================================
if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(INPUT_FILE)

df = pd.read_csv(INPUT_FILE)
date_col = df.columns[0]

adj_cols = [c for c in df.columns if "Adj Close" in c]
if not adj_cols:
    adj_cols = [c for c in df.columns if c != date_col]

prices = df[[date_col] + adj_cols].copy()
# FIX: strip the '_Adj Close' suffix correctly.
prices.columns = ["Date"] + [c.replace("_Adj Close", "").strip() for c in adj_cols]
prices["Date"] = pd.to_datetime(prices["Date"])

# =============================================================================
# STEP 2 : DAILY SIMPLE RETURNS
# =============================================================================
returns = prices.set_index("Date").pct_change().dropna()

# =============================================================================
# STEP 3 : MATRICES
# =============================================================================
cov_matrix  = returns.cov()      # sample covariance (ddof=1)
corr_matrix = returns.corr()

# =============================================================================
# STEP 4 : WRITE WORKBOOK
# =============================================================================
wb = Workbook()
fill = PatternFill("solid", fgColor="4F81BD")
font = Font(color="FFFFFF", bold=True)

def write_matrix(ws, title, matrix, fmt):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    headers = list(matrix.columns)
    for j, h in enumerate(headers, start=2):
        c = ws.cell(row=2, column=j, value=h)
        c.fill = fill; c.font = font
    for i, idx in enumerate(matrix.index, start=3):
        c = ws.cell(row=i, column=1, value=idx)
        c.fill = fill; c.font = font
        for j, h in enumerate(headers, start=2):
            cell = ws.cell(row=i, column=j, value=float(matrix.loc[idx, h]))
            cell.number_format = fmt

cov_ws = wb.active
cov_ws.title = "Covariance Matrix"
write_matrix(cov_ws, "Covariance Matrix", cov_matrix, "0.000000")

corr_ws = wb.create_sheet("Correlation Matrix")
write_matrix(corr_ws, "Correlation Matrix", corr_matrix, "0.0000")

ret_ws = wb.create_sheet("Daily Returns")
ret_ws.append(["Date"] + list(returns.columns))
for row in returns.reset_index().itertuples(index=False):
    ret_ws.append(list(row))

ins = wb.create_sheet("Instructions")
notes = [
    "Purpose",
    "Compute covariance and correlation matrices from daily simple returns.",
    "",
    "Definitions",
    "Covariance  : joint movement (scale-dependent).",
    "Correlation : standardized covariance in [-1, +1].",
    "",
    "Interpretation",
    "High positive correlation -> similar movement -> little diversification.",
    "Low / negative correlation -> diversification benefit.",
    "",
    "Common mistakes",
    "1. Using prices instead of returns.",
    "2. Mixing Close and Adjusted Close.",
    "3. Comparing covariance magnitudes across different scales.",
    "",
    "Exercises",
    "1. Identify the most-correlated pair.",
    "2. Identify the least-correlated pair.",
    "3. Discuss the diversification implications for Script 08.",
]
for i, line in enumerate(notes, start=1):
    ins.cell(row=i, column=1, value=line)

os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
wb.save(OUTPUT_FILE)
print("Workbook created:", OUTPUT_FILE)

# =============================================================================
# MINI TEACHING MODULE
# =============================================================================
# WHY RETURNS (not prices)?
#   Prices are non-stationary and depend on arbitrary price levels; covariance
#   and correlation are only meaningful on returns.
#
# EXTENSIONS
#   Conditional-format the correlation sheet to highlight strong pairs.
#   Build a matplotlib heatmap.
#   Compute rolling 60-day correlation matrices.
