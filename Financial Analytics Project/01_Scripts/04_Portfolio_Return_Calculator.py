#!/usr/bin/env python3
r"""
===============================================================================
Financial Analytics using Python
Script 04 : Portfolio Return Calculator
Author    : Vipul Mehta

PURPOSE
-------
Computes the historical annual return of each stock and the portfolio's annual
return. It is designed to be HUMAN-ERROR-FREE:

  * Stock names are read automatically FROM the price file -- nothing to type,
    so a mistyped ticker can never break the run.
  * The portfolio opens EQUAL-WEIGHTED (1/N) so it is correct out of the box.
  * Weights are then edited INSIDE the Excel file, in colour-coded cells, with a
    live check that they still sum to 100%. Everything recomputes automatically.

    Historical annual return = average daily return x annualisation factor
    Portfolio return         = SUM( weight x historical annual return )

INPUT  : Yahoo_Daily_Prices_2016_2026.csv
OUTPUT : Portfolio_Return_Summary.xlsx

COLOUR CONVENTION (same as our other models)
    BLUE  = inputs you may edit (the weights, the annualisation factor)
    BLACK = formulas on the summary sheet
    GREEN = formulas that read the Calculations (daily returns) sheet
===============================================================================
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

# ========================= USER SETTINGS =====================================
# The ONLY thing to set is the folder. It must match where Script 01 saved the
# CSV. Portfolio weights are NOT set here any more -- you change them inside the Excel file.

FOLDER = r"D:\Trimester 4\Finance\Financial Analytics Project\02_Output_Assignment_1"
ANNUALIZATION_FACTOR = 250          # seeds the (editable) blue cell in the sheet

# =============================================================================

INPUT_FILE  = os.path.join(FOLDER, "Yahoo_Daily_Prices_2016_2026.csv")
OUTPUT_FILE = os.path.join(FOLDER, "Portfolio_Return_Summary.xlsx")

# =============================================================================
# STEP 1 : READ PRICES AND CAPTURE TICKERS FROM THE FILE
# =============================================================================
if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(
        f"Could not find:\n{INPUT_FILE}\n"
        "Run Script 01 first and make sure FOLDER matches where it saved the CSV."
    )

df = pd.read_csv(INPUT_FILE)
date_col = df.columns[0]

# Every stock in the file is identified by its Adjusted Close column.
adj_cols = [c for c in df.columns if "Adj Close" in c]
if not adj_cols:
    raise ValueError("No 'Adj Close' columns found -- re-run Script 01.")

prices = df[[date_col] + adj_cols].copy()
prices.columns = ["Date"] + [c.replace("_Adj Close", "").strip() for c in adj_cols]
prices["Date"] = pd.to_datetime(prices["Date"])

# =============================================================================
# STEP 2 : DAILY SIMPLE RETURNS  (the raw input the sheet reads)
# =============================================================================
returns = prices.set_index("Date").pct_change().dropna()
stocks  = list(returns.columns)
N       = len(stocks)
n_days  = len(returns)
print(f"Stocks captured from file ({N}): {stocks}")
print(f"Starting equal-weighted at {1/N:.1%} each. Edit weights in the Excel file.")

# =============================================================================
# STEP 3 : STYLES
# =============================================================================
BLUE  = Font(name="Arial", color="0000FF")
BLACK = Font(name="Arial", color="000000")
GREEN = Font(name="Arial", color="008000")
BOLD  = Font(name="Arial", bold=True)
WBOLD = Font(name="Arial", bold=True, color="FFFFFF")
TITLE = Font(name="Arial", bold=True, size=14)
HEAD  = PatternFill("solid", fgColor="4F81BD")
YEL   = PatternFill("solid", fgColor="FFF2CC")
REDF  = PatternFill("solid", fgColor="FFC7CE")     # bad total
GRNF  = PatternFill("solid", fgColor="C6EFCE")     # good total
thin  = Side(style="thin", color="BFBFBF")
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR   = Alignment(horizontal="center")
PCT   = "0.00%"

wb = Workbook()

# =============================================================================
# STEP 4 : CALCULATIONS SHEET  (daily returns behind the averages)
# =============================================================================
calc = wb.create_sheet("Calculations")
calc.append(["Date"] + stocks)
for c in calc[1]:
    c.fill = HEAD; c.font = WBOLD
for row in returns.reset_index().itertuples(index=False):
    calc.append([row[0].date()] + [float(v) for v in row[1:]])
for r in calc.iter_rows(min_row=2, min_col=2):
    for c in r:
        c.number_format = "0.0000%"
calc_last = n_days + 1
def calc_range(k):
    col = get_column_letter(2 + k)      # B, C, D ... on the Calculations sheet
    return f"'Calculations'!{col}2:{col}{calc_last}"

# =============================================================================
# STEP 5 : PORTFOLIO SUMMARY SHEET  (all formulas -> fully live)
# =============================================================================
ws = wb.create_sheet("Portfolio Summary", 0)
ws["A1"] = "Portfolio Return Calculator"
ws["A1"].font = TITLE

# Annualisation factor (blue, editable) -- the green formulas below multiply by it.
ws["A2"] = "Annualisation factor (trading days/yr)"; ws["A2"].font = BOLD
ws["B2"] = ANNUALIZATION_FACTOR; ws["B2"].font = BLUE; ws["B2"].border = BOX

HEAD_ROW = 4
AR_ROW   = 5      # Historical Annual Return (green)
W_ROW    = 6      # Weight (blue, editable)
WC_ROW   = 7      # Weighted Contribution (black)
PR_ROW   = 9      # Portfolio Annual Return (black, headline)

first_col   = 2                 # column B = first stock
last_col_L  = get_column_letter(1 + N)      # last stock column letter
total_col   = 2 + N             # one past the last stock
total_col_L = get_column_letter(total_col)

# ---- header row -----------------------------------------------------------
ws.cell(row=HEAD_ROW, column=1, value="Metric").fill = HEAD
ws.cell(row=HEAD_ROW, column=1).font = WBOLD
for k, s in enumerate(stocks):
    c = ws.cell(row=HEAD_ROW, column=2 + k, value=s)
    c.fill = HEAD; c.font = WBOLD; c.alignment = CTR; c.border = BOX
tc = ws.cell(row=HEAD_ROW, column=total_col, value="Total")
tc.fill = HEAD; tc.font = WBOLD; tc.alignment = CTR; tc.border = BOX

# ---- row labels -----------------------------------------------------------
ws.cell(row=AR_ROW, column=1, value="Historical Annual Return").font = BOLD
ws.cell(row=W_ROW,  column=1, value="Portfolio Weight  (edit these)").font = BOLD
ws.cell(row=WC_ROW, column=1, value="Weighted Contribution").font = BOLD

# ---- per-stock cells ------------------------------------------------------
for k in range(N):
    col_L = get_column_letter(2 + k)

    # Historical annual return = AVERAGE(daily returns) x factor  (GREEN)
    ar = ws.cell(row=AR_ROW, column=2 + k, value=f"=AVERAGE({calc_range(k)})*$B$2")
    ar.font = GREEN; ar.number_format = PCT; ar.border = BOX

    # Weight = 1/N default  (BLUE, editable)
    w = ws.cell(row=W_ROW, column=2 + k, value=round(1.0 / N, 10))
    w.font = BLUE; w.number_format = PCT; w.border = BOX

    # Weighted contribution = weight x annual return  (BLACK)
    wc = ws.cell(row=WC_ROW, column=2 + k,
                 value=f"={col_L}{W_ROW}*{col_L}{AR_ROW}")
    wc.font = BLACK; wc.number_format = PCT; wc.border = BOX

# ---- totals column --------------------------------------------------------
# Total weight (should be 100%) with a live OK / FIX check + colour.
tot_w = ws.cell(row=W_ROW, column=total_col,
                value=f"=SUM({get_column_letter(first_col)}{W_ROW}:{last_col_L}{W_ROW})")
tot_w.font = BOLD; tot_w.number_format = PCT; tot_w.border = BOX
ws.cell(row=W_ROW, column=total_col + 1,
        value=f'=IF(ABS({total_col_L}{W_ROW}-1)<0.0001,"OK","FIX: must total 100%")').font = BOLD

# Total weighted contribution = portfolio return (shown again below too).
tot_wc = ws.cell(row=WC_ROW, column=total_col,
                 value=f"=SUM({get_column_letter(first_col)}{WC_ROW}:{last_col_L}{WC_ROW})")
tot_wc.font = BOLD; tot_wc.number_format = PCT; tot_wc.border = BOX

# ---- headline result ------------------------------------------------------
ws.cell(row=PR_ROW, column=1, value="Portfolio Annual Return").font = BOLD
pr = ws.cell(row=PR_ROW, column=2, value=f"={total_col_L}{WC_ROW}")
pr.font = BLACK; pr.number_format = PCT; pr.border = BOX; pr.fill = YEL

# ---- conditional formatting on the weight total (red if not 100%) ---------
tw_ref = f"{total_col_L}{W_ROW}"
ws.conditional_formatting.add(tw_ref,
    FormulaRule(formula=[f"ABS({tw_ref}-1)>0.0001"], fill=REDF))
ws.conditional_formatting.add(tw_ref,
    FormulaRule(formula=[f"ABS({tw_ref}-1)<=0.0001"], fill=GRNF))

# ---- widths ---------------------------------------------------------------
ws.column_dimensions["A"].width = 34
for k in range(N + 2):
    ws.column_dimensions[get_column_letter(2 + k)].width = 13

# =============================================================================
# STEP 6 : INSTRUCTIONS SHEET
# =============================================================================
ins = wb.create_sheet("Instructions")
lines = [
    ("Portfolio Return Calculator", TITLE),
    ("", BLACK),
    ("Colour legend", BOLD),
    ("  BLUE  = inputs you may edit (weights, annualisation factor)", BLUE),
    ("  BLACK = formulas on the Portfolio Summary sheet", BLACK),
    ("  GREEN = formulas that read the Calculations sheet", GREEN),
    ("", BLACK),
    ("How to use", BOLD),
    ("  1. The portfolio opens equal-weighted and already correct.", BLACK),
    ("  2. Edit the BLUE weight cells. Keep the Total at 100% (it turns", BLACK),
    ("     green when OK, red when it does not add up).", BLACK),
    ("  3. Portfolio Annual Return updates automatically.", BLACK),
    ("", BLACK),
    ("Formulae", BOLD),
    ("  Historical Annual Return = Average Daily Return x factor", BLACK),
    ("  Portfolio Return = SUM( Weight x Historical Annual Return )", BLACK),
    ("", BLACK),
    ("Exercises", BOLD),
    ("  1. Tilt toward the highest-return stock; watch the portfolio return rise.", BLACK),
    ("  2. Set one weight to 100% and the rest to 0%.", BLACK),
    ("  3. Change the annualisation factor from 250 to 252.", BLACK),
    ("  4. Confirm the weighted contributions add up to the portfolio return.", BLACK),
]
for i, (t, f) in enumerate(lines, start=1):
    ins.cell(row=i, column=1, value=t).font = f
ins.column_dimensions["A"].width = 66

# =============================================================================
# STEP 7 : SAVE
# =============================================================================
os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
wb.save(OUTPUT_FILE)
print("Workbook created:", OUTPUT_FILE)
