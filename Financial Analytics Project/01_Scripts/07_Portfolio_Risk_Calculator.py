#!/usr/bin/env python3
r"""
===============================================================================
Financial Analytics using Python
Script 08 : Portfolio Risk Calculator  (w' Sigma w)
Author    : Vipul Mehta
===============================================================================

PURPOSE
-------
This is the capstone of the risk block. It reads the daily price file, converts
Adjusted Close prices into daily SIMPLE returns, and builds an Excel workbook in
which the WHOLE portfolio-risk calculation is done with LIVE EXCEL FORMULAS, not
Python-computed numbers. A student can:

    * see the daily return of every stock on every day (the raw input),
    * see each pairwise covariance built from those returns with SUMPRODUCT,
    * see each weighted contribution  w_i * w_j * Cov(i, j),
    * change any weight and watch portfolio variance and volatility recompute.

WHY SIMPLE RETURNS (not log) FOR PORTFOLIO RISK
-----------------------------------------------
A portfolio return is a weighted sum of its constituents' returns. That linear
relationship holds for SIMPLE returns but NOT for log returns (you cannot add
weighted log returns across assets). So covariance-based portfolio variance must
be built on simple returns. This matches scripts 04, 05 and 06.

THE MATH
--------
    Portfolio variance   sigma_p^2 = SUM_i SUM_j  w_i * w_j * Cov(r_i, r_j)
                                   = w' Sigma w
    Portfolio volatility sigma_p   = sqrt(sigma_p^2)
Everything is computed on DAILY covariance first, then annualised by the same
factor used elsewhere in the course (250 trading days by default).

INPUT
-----
Yahoo_Daily_Prices_2016_2026.csv

OUTPUT
------
Portfolio_Risk_Summary.xlsx
    Sheet 1 "Portfolio Risk"  : weights, covariance matrix, contribution matrix,
                                and the portfolio variance / volatility results.
    Sheet 2 "Daily Returns"   : the daily simple returns the formulas read from.
    Sheet 3 "Instructions"    : notes, colour legend and classroom exercises.

COLOUR CONVENTION (same as our Excel models)
--------------------------------------------
    BLUE  text  = hardcoded inputs you may edit (weights, annualisation factor)
    BLACK text  = formulas computed on this sheet
    GREEN text  = formulas that pull from another sheet (the Daily Returns tab)

TEACHING NOTE
-------------
Students edit ONLY the blue cells (the weights, and optionally the annualisation
factor). The workbook opens with EQUAL weights (1/N) so it runs for any number of
stocks with no editing, then invites re-weighting as the first exercise.
===============================================================================
"""

# =============================================================================
# IMPORT LIBRARIES
# =============================================================================
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# =============================================================================
# USER SETTINGS  ->  edit these two lines only
# =============================================================================
FOLDER = r"D:\Trimester 4\Finance\Financial Analytics\Financial Analytics Project\02_Output"
ANNUALIZATION_FACTOR = 250          # 250 trading days; try 252 as an exercise

INPUT_FILE  = os.path.join(FOLDER, "Yahoo_Daily_Prices_2016_2026.csv")
OUTPUT_FILE = os.path.join(FOLDER, "Portfolio_Risk_Summary.xlsx")

# =============================================================================
# STEP 1 : READ PRICES AND EXTRACT ADJUSTED CLOSE
# =============================================================================
if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(
        f"Could not find the price file:\n{INPUT_FILE}\n"
        "Run script 01 first, and make sure FOLDER matches across all scripts."
    )

df = pd.read_csv(INPUT_FILE)
date_col = df.columns[0]

# Adjusted Close columns look like 'RELIANCE.NS_Adj Close'. Fall back to every
# non-date column only if no adjusted-close columns are present.
adj_cols = [c for c in df.columns if "Adj Close" in c] or \
           [c for c in df.columns if c != date_col]
if not adj_cols:
    raise ValueError("No price columns found in the input file.")

prices = df[[date_col] + adj_cols].copy()
# Turn 'RELIANCE.NS_Adj Close' into a clean ticker label 'RELIANCE.NS'.
prices.columns = ["Date"] + [c.replace("_Adj Close", "").strip() for c in adj_cols]
prices["Date"] = pd.to_datetime(prices["Date"])

# =============================================================================
# STEP 2 : DAILY SIMPLE RETURNS
# =============================================================================
# dropna() aligns all stocks to their common set of trading days, so every
# covariance below is computed on the same, complete rows.
returns = prices.set_index("Date").pct_change().dropna()
stocks = list(returns.columns)
N = len(stocks)
n_days = len(returns)
if N < 2:
    raise ValueError("Portfolio risk needs at least two stocks.")

print(f"Stocks: {stocks}")
print(f"Return observations per stock: {n_days}")

# =============================================================================
# STEP 3 : STYLES
# =============================================================================
BLUE  = Font(name="Arial", color="0000FF")            # editable inputs
BLACK = Font(name="Arial", color="000000")            # same-sheet formulas
GREEN = Font(name="Arial", color="008000")            # cross-sheet formulas
BOLD  = Font(name="Arial", bold=True)
WHITE_BOLD = Font(name="Arial", bold=True, color="FFFFFF")
TITLE = Font(name="Arial", bold=True, size=14)

HEAD_FILL = PatternFill("solid", fgColor="4F81BD")
YELLOW    = PatternFill("solid", fgColor="FFF2CC")     # highlight headline result
thin = Side(style="thin", color="BFBFBF")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center")

PCT   = "0.00%"
PCT4  = "0.0000%"
COV_FMT = "0.000000"

wb = Workbook()

# =============================================================================
# STEP 4 : DAILY RETURNS SHEET  (the raw input the formulas read)
# =============================================================================
# Written FIRST so the Portfolio Risk sheet can point its formulas at it.
ret_ws = wb.create_sheet("Daily Returns")
ret_ws.append(["Date"] + stocks)
for c in ret_ws[1]:
    c.fill = HEAD_FILL
    c.font = WHITE_BOLD
for row in returns.reset_index().itertuples(index=False):
    ret_ws.append([row[0].date()] + [float(v) for v in row[1:]])
for r in ret_ws.iter_rows(min_row=2, min_col=2):
    for c in r:
        c.number_format = PCT4

# Row range of the return data on that sheet (row 1 is the header).
ret_first, ret_last = 2, n_days + 1
# Excel range for stock k's return column on the Daily Returns sheet.
def ret_range(k):
    col = get_column_letter(2 + k)          # B, C, D, ...
    return f"'Daily Returns'!{col}{ret_first}:{col}{ret_last}"

# =============================================================================
# STEP 5 : PORTFOLIO RISK SHEET
# =============================================================================
ws = wb.create_sheet("Portfolio Risk", 0)   # make it the first tab
ws["A1"] = "Portfolio Risk   (sigma_p^2 = w' Sigma w)"
ws["A1"].font = TITLE

# ---- 5a. Annualisation factor (blue input) --------------------------------
ws["A3"] = "Annualisation factor (trading days/yr)"
ws["A3"].font = BOLD
fcell = "B3"
ws[fcell] = ANNUALIZATION_FACTOR
ws[fcell].font = BLUE
ws[fcell].border = BOX

# ---- 5b. Weights (blue inputs, default equal) -----------------------------
ws["A5"] = "WEIGHTS  -  edit the blue cells; they must sum to 100%"
ws["A5"].font = BOLD
wlabel_row = 6                              # stock names across this row
wrow = 7                                    # weights live on this row
ws.cell(row=wlabel_row, column=1, value="Stock").font = WHITE_BOLD
ws.cell(row=wlabel_row, column=1).fill = HEAD_FILL
ws.cell(row=wrow, column=1, value="Weight").font = BOLD

weight_cell = {}                            # stock index -> weight cell (e.g. 'B7')
for k, s in enumerate(stocks):
    col = get_column_letter(2 + k)
    hc = ws.cell(row=wlabel_row, column=2 + k, value=s)
    hc.fill = HEAD_FILL; hc.font = WHITE_BOLD; hc.alignment = CENTER; hc.border = BOX
    wc = ws.cell(row=wrow, column=2 + k, value=round(1.0 / N, 10))   # equal weight
    wc.font = BLUE; wc.number_format = PCT; wc.border = BOX
    weight_cell[k] = f"{col}{wrow}"

# Weight total + a visible check that it equals 100%.
tot_col = get_column_letter(2 + N)
ws.cell(row=wlabel_row, column=2 + N, value="Total").font = WHITE_BOLD
ws.cell(row=wlabel_row, column=2 + N).fill = HEAD_FILL
ws.cell(row=wlabel_row, column=2 + N).border = BOX
tcell = ws.cell(row=wrow, column=2 + N,
                value=f"=SUM({weight_cell[0]}:{weight_cell[N-1]})")
tcell.font = BOLD; tcell.number_format = PCT; tcell.border = BOX
ws.cell(row=wrow, column=3 + N,
        value=f'=IF(ABS({tot_col}{wrow}-1)<0.0001,"OK","FIX: must total 100%")').font = BOLD

# ---- 5c. Daily covariance matrix (green cross-sheet formulas) --------------
cov_title = wrow + 2
ws.cell(row=cov_title, column=1,
        value="DAILY COVARIANCE MATRIX  =SUMPRODUCT(...)/(n-1) on daily returns").font = BOLD
cov_head = cov_title + 1                     # column headers of the matrix
cov_first = cov_head + 1                     # first data row
for k, s in enumerate(stocks):               # column headers + row headers
    hc = ws.cell(row=cov_head, column=2 + k, value=s)
    hc.fill = HEAD_FILL; hc.font = WHITE_BOLD; hc.alignment = CENTER; hc.border = BOX
    rc = ws.cell(row=cov_first + k, column=1, value=s)
    rc.fill = HEAD_FILL; rc.font = WHITE_BOLD; rc.border = BOX

def cov_cell(i, j):
    return f"{get_column_letter(2 + j)}{cov_first + i}"

for i in range(N):
    for j in range(N):
        xi, yj = ret_range(i), ret_range(j)
        # Sample covariance (ddof = 1), identical to pandas .cov():
        f = (f"=SUMPRODUCT(({xi}-AVERAGE({xi})),({yj}-AVERAGE({yj})))"
             f"/(COUNT({xi})-1)")
        c = ws[cov_cell(i, j)]
        c.value = f; c.font = GREEN; c.number_format = COV_FMT; c.border = BOX

# ---- 5d. Weighted contribution matrix  w_i * w_j * Cov(i,j) ----------------
con_title = cov_first + N + 1
ws.cell(row=con_title, column=1,
        value="WEIGHTED CONTRIBUTION MATRIX  = w_i * w_j * Cov(i,j)").font = BOLD
con_head = con_title + 1
con_first = con_head + 1
for k, s in enumerate(stocks):
    hc = ws.cell(row=con_head, column=2 + k, value=s)
    hc.fill = HEAD_FILL; hc.font = WHITE_BOLD; hc.alignment = CENTER; hc.border = BOX
    rc = ws.cell(row=con_first + k, column=1, value=s)
    rc.fill = HEAD_FILL; rc.font = WHITE_BOLD; rc.border = BOX

for i in range(N):
    wi_col = get_column_letter(2 + i)        # weight of row-stock i (row wrow)
    for j in range(N):
        wj_col = get_column_letter(2 + j)    # weight of col-stock j (row wrow)
        # Mixed absolute refs so the block is safe to copy/inspect in Excel.
        f = f"=${wi_col}${wrow}*${wj_col}${wrow}*{cov_cell(i, j)}"
        c = ws.cell(row=con_first + i, column=2 + j, value=f)
        c.font = BLACK; c.number_format = COV_FMT; c.border = BOX

con_last = con_first + N - 1
con_left, con_rightcol = "B", get_column_letter(1 + N)

# ---- 5e. Results ----------------------------------------------------------
res = con_last + 2
def result(label, formula, fmt, fill=None, note=None):
    global res
    lc = ws.cell(row=res, column=1, value=label); lc.font = BOLD
    vc = ws.cell(row=res, column=2, value=formula)
    vc.font = BLACK; vc.number_format = fmt; vc.border = BOX
    if fill: vc.fill = fill; lc.fill = fill
    if note:
        ws.cell(row=res, column=3, value=note).font = Font(name="Arial", italic=True)
    res += 1

contrib_block = f"{con_left}{con_first}:{con_rightcol}{con_last}"
result("Daily portfolio variance",
       f"=SUM({contrib_block})", COV_FMT,
       note="Sum of every weighted contribution above")
dvar_row = res - 1
result("Annualised portfolio variance",
       f"=B{dvar_row}*$B$3", COV_FMT,
       note="Daily variance x annualisation factor")
avar_row = res - 1
result("Annualised portfolio volatility (SD)",
       f"=SQRT(B{avar_row})", PCT, fill=YELLOW,
       note="Headline number: portfolio risk")
psd_row = res - 1

# Undiversified benchmark = weighted average of individual annualised SDs.
indiv_sd_terms = "+".join(
    f"{weight_cell[k]}*SQRT({cov_cell(k, k)}*$B$3)" for k in range(N)
)
result("Weighted-avg volatility (no diversification)",
       f"={indiv_sd_terms}", PCT,
       note="What risk would be if correlations were +1")
wavg_row = res - 1
result("Diversification benefit",
       f"=B{wavg_row}-B{psd_row}", PCT,
       note="Weighted-avg SD minus actual portfolio SD (>=0)")

# ---- 5f. Column widths ----------------------------------------------------
ws.column_dimensions["A"].width = 42
for k in range(N + 2):
    ws.column_dimensions[get_column_letter(2 + k)].width = 13

# =============================================================================
# STEP 6 : INSTRUCTIONS SHEET
# =============================================================================
ins = wb.create_sheet("Instructions")
notes = [
 ("Portfolio Risk Calculator", TITLE),
 ("", BLACK),
 ("Colour legend", BOLD),
 ("  BLUE  cells  = inputs you may edit (weights, annualisation factor)", BLUE),
 ("  BLACK cells  = formulas on the Portfolio Risk sheet", BLACK),
 ("  GREEN cells  = formulas that read from the Daily Returns sheet", GREEN),
 ("", BLACK),
 ("How it works", BOLD),
 ("  1. Daily Returns sheet holds daily simple returns (the raw input).", BLACK),
 ("  2. Covariance matrix = SUMPRODUCT(dev_i, dev_j)/(n-1) on those returns.", BLACK),
 ("  3. Contribution matrix cell = w_i * w_j * Cov(i,j).", BLACK),
 ("  4. Daily portfolio variance = SUM of the contribution matrix = w' Sigma w.", BLACK),
 ("  5. Annualise (x trading days), then take the square root for volatility.", BLACK),
 ("", BLACK),
 ("Why simple (not log) returns", BOLD),
 ("  Portfolio return is a weighted SUM of asset returns; that is linear for", BLACK),
 ("  simple returns only. Log returns cannot be aggregated across assets.", BLACK),
 ("", BLACK),
 ("Exercises", BOLD),
 ("  1. Start at equal weights, then tilt toward the lowest-volatility stock.", BLACK),
 ("     Watch portfolio SD fall below the weighted-average SD.", BLACK),
 ("  2. Put 100% in one stock; confirm portfolio SD = that stock's annual SD.", BLACK),
 ("  3. Find weights (they must sum to 100%) that minimise portfolio SD by hand.", BLACK),
 ("  4. Switch the annualisation factor from 250 to 252 and note the change.", BLACK),
 ("  5. Explain why diversification benefit grows as correlations fall.", BLACK),
]
for i, (line, font) in enumerate(notes, start=1):
    c = ins.cell(row=i, column=1, value=line)
    c.font = font
ins.column_dimensions["A"].width = 74

# =============================================================================
# STEP 7 : SAVE
# =============================================================================
os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
wb.save(OUTPUT_FILE)
print("Portfolio risk workbook created:", OUTPUT_FILE)
