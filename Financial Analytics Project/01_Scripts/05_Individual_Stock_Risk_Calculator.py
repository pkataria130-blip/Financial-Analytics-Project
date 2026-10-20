#!/usr/bin/env python3
r"""
===============================================================================
Financial Analytics using Python
Script 05 : Individual Stock Risk Calculator
Author    : Vipul Mehta
===============================================================================

PURPOSE
-------
Computes the historical risk of each stock from daily SIMPLE returns derived
from Adjusted Close prices.

MEASURES
--------
1. Daily variance of returns
2. Annualized variance          = daily variance x 250
3. Annualized standard deviation = sqrt(annualized variance)   [volatility]
4. Annualized return            = average daily return x 250
5. Coefficient of variation      = annualized SD / annualized return   [risk per unit return]

INPUT  : Yahoo_Daily_Prices_2016_2026.csv
OUTPUT : Risk_Summary.xlsx

ASSUMPTIONS
-----------
Uses Adjusted Close, 250 trading days/year, risk from daily simple returns.

Modify only INPUT_FILE / FOLDER or ANNUALIZATION_FACTOR. Everything
updates automatically for any number of stocks.
===============================================================================
"""

import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# =============================================================================
# SETTINGS
# =============================================================================
FOLDER = r"D:\Trimester 4\Finance\Financial Analytics\Financial Analytics Project\02_Output"
ANNUALIZATION_FACTOR = 250

INPUT_FILE  = os.path.join(FOLDER, "Yahoo_Daily_Prices_2016_2026.csv")
OUTPUT_FILE = os.path.join(FOLDER, "Risk_Summary.xlsx")

# =============================================================================
# STEP 1 : READ PRICES
# =============================================================================
if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(INPUT_FILE)

df = pd.read_csv(INPUT_FILE)
date_col = df.columns[0]

# Identify Adjusted Close columns (fall back to all non-date columns).
adj_cols = [c for c in df.columns if "Adj Close" in c] or \
           [c for c in df.columns if c != date_col]

prices = df[[date_col] + adj_cols].copy()
# FIX: strip the '_Adj Close' suffix correctly (old code used a trailing space
# that never matched, so labels kept the ugly '_Adj Close' text).
prices.columns = ["Date"] + [c.replace("_Adj Close", "").strip() for c in adj_cols]

# =============================================================================
# STEP 2 : DAILY SIMPLE RETURNS
# =============================================================================
returns = prices.set_index("Date").pct_change().dropna()

# =============================================================================
# STEP 3 : RISK MEASURES
# =============================================================================
daily_var     = returns.var()                              # sample variance (ddof=1)
annual_var    = daily_var * ANNUALIZATION_FACTOR
annual_sd     = np.sqrt(annual_var)
annual_return = returns.mean() * ANNUALIZATION_FACTOR
cv            = annual_sd / annual_return.replace(0, np.nan)

# =============================================================================
# STEP 4 : SUMMARY SHEET
# =============================================================================
wb = Workbook()
ws = wb.active
ws.title = "Risk Summary"

header_fill = PatternFill("solid", fgColor="4F81BD")
header_font = Font(color="FFFFFF", bold=True)

headers = ["Measure"] + list(returns.columns)
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill; cell.font = header_font

metrics = [
    ("Daily Variance",                daily_var,     "0.000000"),
    ("Annualized Variance",           annual_var,    "0.000000"),
    ("Annualized Standard Deviation", annual_sd,     "0.00%"),
    ("Annualized Return",             annual_return, "0.00%"),
    ("Coefficient of Variation (CV)", cv,            "0.0000"),
]
for r, (label, values, fmt) in enumerate(metrics, start=2):
    ws.cell(row=r, column=1, value=label)
    for c, stock in enumerate(returns.columns, start=2):
        cell = ws.cell(row=r, column=c)
        cell.value = float(values[stock]) if pd.notna(values[stock]) else None
        cell.number_format = fmt

# =============================================================================
# STEP 5 : DAILY RETURNS SHEET
# =============================================================================
calc = wb.create_sheet("Daily Returns")
calc.append(["Date"] + list(returns.columns))
for row in returns.reset_index().itertuples(index=False):
    calc.append(list(row))

# =============================================================================
# STEP 6 : INSTRUCTIONS SHEET
# =============================================================================
ins = wb.create_sheet("Instructions")
notes = [
    "Purpose: Compute historical stock risk.",
    "",
    "Formulae",
    "Daily Variance = VAR(Daily Returns)",
    "Annualized Variance = Daily Variance x 250",
    "Annualized Std Dev = SQRT(Annualized Variance)",
    "Annualized Return = Average Daily Return x 250",
    "CV = Annualized Std Dev / Annualized Return",
    "",
    "Exercises",
    "1. Rank stocks by volatility.",
    "2. Rank stocks by CV.",
    "3. Compare results using 252 trading days.",
    "4. Add more stocks and rerun.",
]
for i, line in enumerate(notes, start=1):
    ins.cell(row=i, column=1, value=line)

os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
wb.save(OUTPUT_FILE)
print("Risk summary workbook created:", OUTPUT_FILE)

# =============================================================================
# MINI TEACHING MODULE
# =============================================================================
# INTERPRETATION
#   Higher variance / SD  = more uncertainty (volatility).
#   Lower CV              = better risk-return trade-off.
#
# COMMON STUDENT MISTAKES
#   Computing variance from prices instead of returns.
#   Forgetting to annualize.
#   Using Close instead of Adjusted Close.
#   Mixing decimal and percentage formats.
